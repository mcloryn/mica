"""Pembuat file contoh kecil untuk tes M-02 (TXT, MD, PDF, DOCX)."""

from pathlib import Path

import docx
from pypdf import PdfWriter

TXT = "Universitas Mikroskil membuka pendaftaran mahasiswa baru.\n\nPendaftaran dibuka setiap hari kerja."

MD = """Pengantar sebelum heading.

# Biaya

Biaya pendaftaran Rp300.000.

| Prodi | Biaya per semester |
| --- | --- |
| Informatika | Rp7.000.000 |
| Sistem Informasi | Rp6.500.000 |

## Beasiswa

Beasiswa prestasi tersedia untuk nilai rapor minimal 85.
"""


def write_pdf_pages(path: Path, streams: list[str]) -> None:
    """PDF Helvetica dengan satu content stream per halaman, ditulis manual (tanpa library)."""
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"",  # /Pages, diisi setelah semua halaman dibuat
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    kids = []
    for stream in streams:
        data = stream.encode("latin-1")
        objects.append(b"<< /Length %d >>\nstream\n" % len(data) + data + b"\nendstream")
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % len(objects)
        )
        kids.append(len(objects))
    refs = " ".join(f"{k} 0 R" for k in kids)
    objects[1] = f"<< /Type /Pages /Kids [{refs}] /Count {len(kids)} >>".encode()
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    path.write_bytes(bytes(out))


def _escape(line: str) -> str:
    return line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def lines_stream(lines: list[str]) -> str:
    """Content stream satu halaman: satu baris per Tj, jarak baris 14 pt."""
    ops = ["BT", "/F1 12 Tf", "72 720 Td", "14 TL"]
    ops += [f"({_escape(line)}) Tj T*" for line in lines]
    ops.append("ET")
    return "\n".join(ops)


def write_text_pdf(path: Path, lines: list[str]) -> None:
    """PDF satu halaman berteks."""
    write_pdf_pages(path, [lines_stream(lines)])


# lebar glyph Helvetica (AFM, per 1000 em) untuk huruf yang dipakai split_word_stream
_HELVETICA = {"s": 500, "e": 556, "m": 833, "t": 278, "r": 333, " ": 278, "d": 556, "a": 556}


def split_word_stream(left: str, right: str, kern: int = 100) -> str:
    """Satu baris yang dipecah menjadi dua objek teks di tengah kata, seperti PDF rata kanan-kiri.

    Objek pertama memakai geseran TJ -kern di antara hurufnya; objek kedua mulai tepat di ujung
    objek pertama (posisi mutlak Tm). kern=0 = tanpa geseran (pembanding).
    """
    width = sum(_HELVETICA[c] for c in left) + kern * (len(left) - 1)
    x = 72 + width * 12 / 1000
    tj = f" -{kern} ".join(f"({c})" for c in left) if kern else f"({left})"
    return (
        f"BT /F1 12 Tf 1 0 0 1 72 720 Tm [{tj}] TJ ET\n"
        f"BT /F1 12 Tf 1 0 0 1 {x:.3f} 720 Tm [({right})] TJ ET"
    )


def write_blank_pdf(path: Path) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with path.open("wb") as f:
        writer.write(f)


def write_docx(path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Paragraf pembuka.")
    document.add_heading("Jalur Masuk", level=1)
    document.add_paragraph("Jalur undangan untuk siswa berprestasi.")
    table = document.add_table(rows=3, cols=2)
    for r, row in enumerate([("Jalur", "Syarat"), ("Undangan", "Rapor"), ("Reguler", "Tes")]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    document.add_heading("Fasilitas", level=2)
    document.add_paragraph("Laboratorium komputer tersedia.")
    document.save(str(path))


SOURCES_YAML = """\
brosur.txt: {title: "Brosur PMB", valid_from: 2026-08-01}
biaya.md: {title: "Rincian Biaya", valid_from: 2026-08-01}
panduan.pdf: {title: "Panduan Pendaftaran", valid_from: "2026-07-15"}
jalur.docx: {title: "Jalur Masuk", valid_from: 2026-09-01}
"""


def make_kb_folder(folder: Path, sources_yaml: str = SOURCES_YAML) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "brosur.txt").write_text(TXT, encoding="utf-8")
    (folder / "biaya.md").write_text(MD, encoding="utf-8")
    write_text_pdf(folder / "panduan.pdf", ["Panduan pendaftaran online.", "Unggah rapor (PDF)."])
    write_docx(folder / "jalur.docx")
    (folder / "sources.yaml").write_text(sources_yaml, encoding="utf-8")
    return folder


class FakeEmbedder:
    """Embedder deterministik untuk tes: vektor dari hash teks, dinormalisasi; token = kata."""

    model_id = "fake/embedder"
    revision = "0" * 40
    weights_file = "fake.safetensors"

    def __init__(self, dim: int = 8, fail_on_call: bool = False) -> None:
        self.dim = dim
        self.fail_on_call = fail_on_call
        self.calls = 0

    def embed(self, texts: list[str]):
        import hashlib

        import numpy as np

        self.calls += 1
        if self.fail_on_call:
            raise RuntimeError("embedder gagal (disengaja)")
        rows = []
        for text in texts:
            seed = int.from_bytes(hashlib.sha256(text.encode("utf-8")).digest()[:8], "little")
            rows.append(np.random.default_rng(seed).standard_normal(self.dim))
        matrix = np.asarray(rows, dtype=np.float32)
        return matrix / np.linalg.norm(matrix, axis=1, keepdims=True)

    def count_tokens(self, text: str) -> int:
        return len(text.split())
