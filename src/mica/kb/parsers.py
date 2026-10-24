"""M-02: membaca dokumen sumber (TXT, MD, PDF, DOCX) menjadi ParsedDocument."""

from __future__ import annotations

import dataclasses
import itertools
import logging
import re
from collections.abc import Iterable, Mapping
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from types import MappingProxyType
from typing import Any, Literal, Protocol

import yaml

from mica.core.types import SECTION_SEP, ErrorCode

logger = logging.getLogger(__name__)

SourceType = Literal["official", "dev"]
SOURCES_FILE = "sources.yaml"


class KbBuildError(Exception):
    def __init__(self, code: ErrorCode, message: str) -> None:
        super().__init__(f"{code}: {message}")
        self.code = ErrorCode(code)
        self.message = message


@dataclass(frozen=True)
class Block:
    section: str  # judul bagian; "" bila sebelum heading pertama atau format tanpa heading
    text: str | None = None
    table: tuple[tuple[str, ...], ...] | None = None  # baris pertama = judul kolom
    # nomor halaman PDF (mulai 1) untuk setiap baris text; None untuk format tanpa halaman
    line_pages: tuple[int, ...] | None = None
    # rentang halaman PDF (awal, akhir) seluruh blok; untuk halaman pengganti (page_overrides)
    page_range: tuple[int, int] | None = None

    def __post_init__(self) -> None:
        if (self.text is None) == (self.table is None):
            raise ValueError("Block harus berisi tepat satu dari text atau table")
        if self.line_pages is not None and (
            self.text is None or len(self.line_pages) != len(self.text.split("\n"))
        ):
            raise ValueError("line_pages harus berisi satu nomor halaman per baris text")
        if self.line_pages is not None and self.page_range is not None:
            raise ValueError("Block hanya boleh punya salah satu dari line_pages atau page_range")


@dataclass(frozen=True)
class SupersededSection:
    section: str  # misalnya "Pasal 4" atau "BAB II … › Pasal 4"
    superseded_by: str  # dokumen dan bagian pengganti
    reason: str


@dataclass(frozen=True)
class ParsedDocument:
    doc_id: str
    doc_title: str
    valid_from: date | None
    source_type: SourceType
    blocks: tuple[Block, ...]
    group: str = ""  # subfolder teratas di folder sumber; "" bila file langsung di sana
    # metadata frontmatter Markdown (title, url, date) untuk init-sources; tidak dipakai build
    frontmatter: Mapping[str, str] = field(default_factory=lambda: MappingProxyType({}))
    # halaman PDF yang diganti ketikan ulang: "8-9" -> path relatif di folder sumber (manifest)
    page_overrides: Mapping[str, str] = field(default_factory=lambda: MappingProxyType({}))
    # bagian yang sudah digantikan dokumen lain; chunk-nya tidak masuk indeks (Langkah 64)
    superseded_sections: tuple[SupersededSection, ...] = ()


class DocumentParser(Protocol):
    def parse(self, path: Path) -> ParsedDocument | None:
        """None bila dokumen dilewati (misalnya PDF tanpa teks)."""
        ...


def _document(path: Path, blocks: list[Block]) -> ParsedDocument:
    # Metadata sementara; load_sources mengisinya dari sources.yaml.
    # source_type "dev" sebagai default supaya dokumen tak pernah dianggap resmi tanpa sengaja.
    return ParsedDocument(
        doc_id=path.stem,
        doc_title=path.stem,
        valid_from=None,
        source_type="dev",
        blocks=tuple(blocks),
    )


def _text_block(section: str, lines: list[str]) -> Block | None:
    text = "\n".join(lines).strip()
    return Block(section=section, text=text) if text else None


class TxtParser:
    def parse(self, path: Path) -> ParsedDocument | None:
        block = _text_block("", path.read_text(encoding="utf-8").splitlines())
        return _document(path, [block] if block else [])


_MD_HEADING = re.compile(r"^#{1,6}\s+(.*?)\s*#*\s*$")
_MD_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?\s*$")


FRONTMATTER_KEYS = ("title", "url", "date")


def _split_frontmatter(lines: list[str]) -> tuple[dict[str, str], list[str]]:
    """Pisahkan frontmatter YAML (--- di baris pertama) dari isi Markdown."""
    if not lines or lines[0].strip() != "---":
        return {}, lines
    for end in range(1, len(lines)):
        if lines[end].strip() in ("---", "..."):
            break
    else:
        return {}, lines  # tanpa penutup: bukan frontmatter
    try:
        raw = yaml.safe_load("\n".join(lines[1:end])) or {}
    except yaml.YAMLError:
        raw = {}
    meta = {}
    if isinstance(raw, dict):
        for key in FRONTMATTER_KEYS:
            value = raw.get(key)
            if value is not None and str(value).strip():
                meta[key] = value.isoformat() if isinstance(value, date) else str(value).strip()
    return meta, lines[end + 1 :]


def _md_row(line: str) -> tuple[str, ...]:
    return tuple(cell.strip() for cell in line.strip().strip("|").split("|"))


_HTML_COMMENT = re.compile(r"<!--.*?-->", re.DOTALL)


def strip_html_comments(text: str) -> str:
    """Buang komentar HTML <!-- ... --> (boleh lebih dari satu baris): catatan penyunting,
    bukan isi dokumen. Baris yang tersisa kosong ikut menjadi baris kosong."""
    return _HTML_COMMENT.sub("", text)


def read_markdown(path: Path) -> tuple[dict[str, str], list[str]]:
    """(frontmatter, baris isi) file Markdown, tanpa komentar HTML."""
    return _split_frontmatter(strip_html_comments(path.read_text(encoding="utf-8")).splitlines())


def md_blocks(lines: list[str], nested: bool = False) -> list[Block]:
    """Block dari baris Markdown.

    nested False: section = judul heading terdekat (cara dokumen .md biasa).
    nested True: section = jalur heading bertingkat, misalnya "BAB II … › 2.1 Sistem Rekrutmen"
    (dipakai halaman pengganti PDF, supaya sama dengan jalur bagian PDF).
    """
    blocks: list[Block] = []
    section = ""
    headings: list[tuple[int, str]] = []  # (level, judul) untuk nested
    buffer: list[str] = []
    in_fence = False

    def flush() -> None:
        block = _text_block(section, buffer)
        if block:
            blocks.append(block)
        buffer.clear()

    i = 0
    while i < len(lines):
        line = lines[i]
        if line.lstrip().startswith("```"):
            in_fence = not in_fence
        elif not in_fence and (heading := _MD_HEADING.match(line)):
            flush()
            if nested:
                level = len(line) - len(line.lstrip("#"))
                headings = [h for h in headings if h[0] < level] + [(level, heading.group(1))]
                section = SECTION_SEP.join(title for _, title in headings)
            else:
                section = heading.group(1)
            i += 1
            continue
        elif (
            not in_fence
            and line.lstrip().startswith("|")
            and i + 1 < len(lines)
            and _MD_TABLE_SEP.match(lines[i + 1])
        ):
            # tabel pipa Markdown: baris judul, baris pemisah, lalu baris data
            flush()
            rows = [_md_row(line)]
            i += 2
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(_md_row(lines[i]))
                i += 1
            blocks.append(Block(section=section, table=tuple(rows)))
            continue
        buffer.append(line)
        i += 1
    flush()
    return blocks


class MdParser:
    def parse(self, path: Path) -> ParsedDocument | None:
        frontmatter, lines = read_markdown(path)
        return dataclasses.replace(
            _document(path, md_blocks(lines)), frontmatter=MappingProxyType(frontmatter)
        )


# --- PDF (Langkah 62) ---
#
# Penyebab kata terpotong (diperiksa di Panduan Kemahasiswaan 2021 dan Peraturan Akademik 2025):
# penulis PDF memecah satu baris rata kanan-kiri menjadi beberapa objek teks (BT..ET) berposisi
# mutlak, kadang di tengah kata ("semester d" | "alam"). extract_text() biasa tidak menambahkan
# geseran angka di dalam TJ ke lebar teks, sehingga mengira ada celah di sambungan dan
# menyisipkan spasi. Mode layout menghitung posisi huruf termasuk geseran itu, jadi dipakai
# sebagai bukti: spasi di teks biasa dihapus hanya bila gabungannya muncul utuh di teks layout
# halaman yang sama dan pasangan terpisahnya tidak. Urutan teks tetap dari extract_text() biasa
# (di tabel dua kolom, mode layout mencampur isi sel kiri dan kanan dalam satu baris).

PDF_EXTRACTION = (
    "L62: extract_text + perbaikan spasi (bukti layout) + tanda hubung; heading BAB/Pasal/n.n"
)
_MAX_SPLIT_PARTS = 4  # satu kata paling banyak terpecah menjadi 4 potongan
_HEADING_MAX_CHARS = 100
_DOT_LEADER = re.compile(r"\.{4,}|…")  # baris daftar isi
_BAB = re.compile(r"BAB\s+[IVXLCDM]+")
_PASAL = re.compile(r"Pasal\s+\d+[A-Z]?")
_NUMBERED = re.compile(r"\d+\.\d+\.?\s+[A-Z].*")
_WORD = re.compile(r"\w+(?:-\w+)*")
_LINE_END_HYPHEN = re.compile(r"(\w+)-$")
_LEADING_WORD = re.compile(r"\w+")
_PRINTED_PAGE_NUMBER = re.compile(r"\d{1,3}|[ivxlc]{1,6}")

PdfLine = tuple[str, int]  # (teks baris, nomor halaman PDF mulai 1)


def layout_evidence(layout_text: str) -> tuple[set[str], set[tuple[str, str]]]:
    """Token dan pasangan token bertetangga (juga lintas baris) dari teks mode layout."""
    tokens: set[str] = set()
    pairs: set[tuple[str, str]] = set()
    previous: str | None = None
    for line in layout_text.splitlines():
        words = line.split()
        if not words:
            continue
        tokens.update(words)
        if previous is not None:
            pairs.add((previous, words[0]))
        pairs.update(itertools.pairwise(words))
        previous = words[-1]
    return tokens, pairs


def repair_split_words(line: str, tokens: set[str], pairs: set[tuple[str, str]]) -> str:
    """Hapus spasi palsu di tengah kata; spasi berlebih dirapikan menjadi satu spasi."""
    words = line.split()
    out: list[str] = []
    i = 0
    while i < len(words):
        best = i
        for j in range(i + 1, min(len(words), i + _MAX_SPLIT_PARTS)):
            if (words[j - 1], words[j]) in pairs:
                break  # spasi ini juga ada di layout: spasi asli
            if "".join(words[i : j + 1]) in tokens:
                best = j
        out.append("".join(words[i : best + 1]))
        i = best + 1
    return " ".join(out)


def join_line_end_hyphens(lines: list[PdfLine]) -> list[PdfLine]:
    """Sambung kata yang terpenggal tanda hubung di akhir baris, hanya di halaman yang sama.

    - kata ulang ("masing-" + "masing") atau bentuk bertanda hubung yang ada di dokumen
      ("perundang-" + "undangan"): disambung, tanda hubung tetap;
    - gabungan tanpa tanda hubung ada di dokumen ("melak-" + "sanakan" -> "melaksanakan"):
      tanda hubung dibuang;
    - selain itu dibiarkan apa adanya.
    """
    vocab = {w.lower() for text, _ in lines for w in _WORD.findall(text)}
    out = list(lines)
    emptied: set[int] = set()  # baris yang seluruh isinya pindah ke baris sebelumnya
    for i in range(len(out) - 1):
        text, page = out[i]
        next_text, next_page = out[i + 1]
        tail = _LINE_END_HYPHEN.search(text)
        lead = _LEADING_WORD.match(next_text)
        if tail is None or lead is None or page != next_page:
            continue
        left, right = tail.group(1).lower(), lead.group(0).lower()
        first, _, rest = next_text.partition(" ")
        if left == right or f"{left}-{right}" in vocab:
            joined = text + first
        elif left + right in vocab:
            joined = text[:-1] + first
        else:
            continue
        out[i] = (joined, page)
        out[i + 1] = (rest, next_page)
        if not rest:
            emptied.add(i + 1)  # jangan sampai menjadi baris kosong (= batas paragraf)
    return [line for i, line in enumerate(out) if i not in emptied]


def drop_printed_page_number(lines: list[str]) -> list[str]:
    """Buang nomor halaman cetak ("4", "ii"): baris berisi angka atau romawi kecil saja yang
    menjadi baris pertama atau terakhir halaman. Tanpa ini, nomor di atas halaman berikutnya
    ikut ke bagian sebelumnya dan rentang halaman chunk melebar satu halaman."""
    out = list(lines)
    filled = [i for i, line in enumerate(out) if line.strip()]
    for i in sorted({filled[0], filled[-1]} if filled else set(), reverse=True):
        if _PRINTED_PAGE_NUMBER.fullmatch(out[i].strip()):
            del out[i]
    return out


def pdf_lines(pages: Iterable[tuple[str, str]]) -> list[PdfLine]:
    """Baris bersih dari pasangan (teks biasa, teks layout) per halaman, berurutan."""
    lines: list[PdfLine] = []
    for number, (plain, layout) in enumerate(pages, start=1):
        tokens, pairs = layout_evidence(layout)
        repaired = [repair_split_words(line, tokens, pairs) for line in plain.splitlines()]
        lines.extend((line, number) for line in drop_printed_page_number(repaired))
    return join_line_end_hyphens(lines)


def heading_level(line: str) -> int | None:
    """1 = "BAB <romawi>", 2 = "Pasal <angka>" atau "<angka>.<angka> Judul"; None = bukan heading."""
    line = line.strip()
    if len(line) > _HEADING_MAX_CHARS or _DOT_LEADER.search(line):
        return None
    if _BAB.fullmatch(line):
        return 1
    if _PASAL.fullmatch(line) or _NUMBERED.fullmatch(line):
        return 2
    return None


def _heading_title(line: str) -> str | None:
    """Baris judul di bawah "BAB II"/"Pasal 39": pendek, bukan kalimat, bukan heading."""
    title = line.strip()
    if (
        not title
        or len(title) > _HEADING_MAX_CHARS
        or title[-1] in ".,;:"
        or heading_level(title) is not None
    ):
        return None
    return title


SectionState = tuple[str, str]  # (BAB terakhir, jalur bagian saat ini)


def pdf_blocks(lines: list[PdfLine]) -> list[Block]:
    """Satu Block per bagian, sehingga chunk tidak melewati batas BAB/Pasal/sub-bab.

    section = jalur bagian terdekat, misalnya "BAB III <judul> › 3.1 Registrasi Administrasi".
    Baris heading (dan judul di bawah BAB/Pasal) masuk ke section, tidak ke text.
    """
    return pdf_blocks_from(lines, ("", ""))[0]


def pdf_blocks_from(lines: list[PdfLine], state: SectionState) -> tuple[list[Block], SectionState]:
    """pdf_blocks yang mulai dari jalur bagian state (misalnya sesudah halaman pengganti) dan
    mengembalikan jalur bagian terakhir."""
    blocks: list[Block] = []
    chapter, section = state
    body: list[PdfLine] = []

    def flush() -> None:
        while body and not body[0][0].strip():
            body.pop(0)
        while body and not body[-1][0].strip():
            body.pop()
        if body:
            blocks.append(
                Block(
                    section=section,
                    text="\n".join(text for text, _ in body),
                    line_pages=tuple(page for _, page in body),
                )
            )
        body.clear()

    i = 0
    while i < len(lines):
        level = heading_level(lines[i][0])
        if level is None:
            body.append(lines[i])
            i += 1
            continue
        flush()
        label = lines[i][0].strip()
        i += 1
        if not _NUMBERED.fullmatch(label) and i < len(lines):
            title = _heading_title(lines[i][0])
            if title is not None:
                label = f"{label} {title}"
                i += 1
        if level == 1:
            chapter = section = label
        else:
            section = SECTION_SEP.join(p for p in (chapter, label) if p)
    flush()
    return blocks, (chapter, section)


# --- halaman pengganti (Langkah 63) ---
#
# sources.yaml: page_overrides: {"8-9": "_overrides/pedoman-pmb-2027-hlm8-9.md"} pada entri PDF.
# Teks halaman itu diganti isi .md (tanpa frontmatter dan komentar HTML); heading .md menjadi
# jalur bagian, dan chunk tetap milik dokumen PDF dengan halaman awal-akhir rentang itu.

OVERRIDES_DIR = "_overrides"  # dilewati iter_source_files: bukan dokumen tersendiri
_PAGE_RANGE = re.compile(r"\s*(\d+)\s*(?:-\s*(\d+)\s*)?")

PageOverrides = Mapping[tuple[int, int], Path]


def parse_page_range(key: str) -> tuple[int, int]:
    """ "8-9" -> (8, 9); "8" -> (8, 8)."""
    match = _PAGE_RANGE.fullmatch(str(key))
    if match is None:
        raise ValueError(f'rentang halaman tidak valid: {key!r} (contoh: "8-9" atau "8")')
    start = int(match.group(1))
    end = int(match.group(2) or start)
    if not 1 <= start <= end:
        raise ValueError(f"rentang halaman tidak valid: {key!r}")
    return start, end


def override_blocks(path: Path, pages: tuple[int, int]) -> list[Block]:
    """Block dari ketikan ulang halaman PDF, semuanya dengan page_range = pages."""
    _, lines = read_markdown(path)
    return [dataclasses.replace(b, page_range=pages) for b in md_blocks(lines, nested=True)]


def _state_after(section: str, state: SectionState) -> SectionState:
    """Jalur bagian yang berlanjut ke halaman PDF sesudah halaman pengganti."""
    if not section:
        return state
    first = section.split(SECTION_SEP)[0]
    chapter = first if _BAB.match(first) else state[0]
    return chapter, section


def _check_overrides(path: Path, overrides: PageOverrides, n_pages: int) -> None:
    previous_end = 0
    for (start, end), md in sorted(overrides.items()):
        if end > n_pages:
            raise KbBuildError(
                ErrorCode.KB_INVALID, f"{path}: page_overrides {start}-{end} > {n_pages} halaman"
            )
        if start <= previous_end:
            raise KbBuildError(ErrorCode.KB_INVALID, f"{path}: page_overrides tumpang tindih")
        if not md.is_file():
            raise KbBuildError(ErrorCode.KB_INVALID, f"{path}: file page_overrides {md} tidak ada")
        previous_end = end


def blocks_with_overrides(lines: list[PdfLine], overrides: PageOverrides) -> list[Block]:
    """Block PDF dengan halaman pengganti disisipkan di posisinya (baris halaman itu sudah tidak
    ada di lines)."""
    blocks: list[Block] = []
    state: SectionState = ("", "")
    rest = lines
    for (start, end), md in sorted(overrides.items()):
        before = [line for line in rest if line[1] < start]
        rest = [line for line in rest if line[1] > end]
        new, state = pdf_blocks_from(before, state)
        replaced = override_blocks(md, (start, end))
        blocks += new + replaced
        if replaced:
            state = _state_after(replaced[-1].section, state)
    blocks += pdf_blocks_from(rest, state)[0]
    return blocks


class PdfParser:
    def parse(
        self, path: Path, page_overrides: PageOverrides | None = None
    ) -> ParsedDocument | None:
        from pypdf import PdfReader

        overrides = dict(page_overrides or {})
        reader = PdfReader(path)
        _check_overrides(path, overrides, len(reader.pages))
        replaced = {p for start, end in overrides for p in range(start, end + 1)}
        pages: list[tuple[str, str]] = []
        for number, page in enumerate(reader.pages, start=1):
            if number in replaced:
                pages.append(("", ""))  # isinya dari halaman pengganti
                continue
            plain = page.extract_text() or ""
            # halaman tanpa teks (atau tanpa /Contents) tidak perlu dan tidak bisa dibaca mode layout
            layout = page.extract_text(extraction_mode="layout") if plain.strip() else ""
            pages.append((plain, layout))
        blocks = blocks_with_overrides(pdf_lines(pages), overrides)
        if not blocks:
            # PDF hasil pindai butuh OCR, di luar cakupan (SDD 2.3)
            logger.warning("%s: %s tidak berisi teks, dilewati", ErrorCode.KB_PDF_NO_TEXT, path)
            return None
        return _document(path, blocks)


class DocxParser:
    def parse(self, path: Path) -> ParsedDocument | None:
        import docx
        from docx.text.paragraph import Paragraph

        document = docx.Document(str(path))
        blocks: list[Block] = []
        section = ""
        buffer: list[str] = []

        def flush() -> None:
            block = _text_block(section, buffer)
            if block:
                blocks.append(block)
            buffer.clear()

        # Paragraf dan tabel dibaca berurutan dari body supaya tabel tetap di bagiannya.
        tables = {t._tbl: t for t in document.tables}
        for child in document.element.body.iterchildren():
            if child in tables:
                flush()
                rows = tuple(
                    tuple(cell.text.strip() for cell in row.cells) for row in tables[child].rows
                )
                if rows:
                    blocks.append(Block(section=section, table=rows))
            elif child.tag.endswith("}p"):
                paragraph = Paragraph(child, document)
                style = paragraph.style.name if paragraph.style is not None else ""
                if style.startswith("Heading") and paragraph.text.strip():
                    flush()
                    section = paragraph.text.strip()
                else:
                    buffer.append(paragraph.text)
        flush()
        return _document(path, blocks)


PARSERS: dict[str, DocumentParser] = {
    ".txt": TxtParser(),
    ".md": MdParser(),
    ".pdf": PdfParser(),
    ".docx": DocxParser(),
}


def _read_sources(folder: Path) -> dict[str, dict[str, Any]]:
    path = folder / SOURCES_FILE
    if not path.exists():
        return {}
    raw = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    if not isinstance(raw, dict):
        raise KbBuildError(ErrorCode.KB_INVALID, f"{path}: isi harus mapping nama-file -> metadata")
    return {str(name): (meta or {}) for name, meta in raw.items()}


def _parse_date(value: Any) -> date | None:
    if value is None or isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value))
    except ValueError:
        return None


def iter_source_files(folder: Path) -> list[Path]:
    """Semua file berformat didukung di folder sumber dan subfoldernya, terurut.

    File dan folder yang namanya diawali titik dilewati, juga folder _overrides (halaman
    pengganti PDF, bukan dokumen tersendiri).
    """
    folder = Path(folder)
    files = []
    for path in folder.rglob("*"):
        rel = path.relative_to(folder)
        if any(part.startswith(".") or part == OVERRIDES_DIR for part in rel.parts):
            continue
        if path.is_file() and path.suffix.lower() in PARSERS:
            files.append(path)
    return sorted(files, key=lambda p: p.relative_to(folder).as_posix())


def source_key(folder: Path, path: Path) -> str:
    """Kunci di sources.yaml: path relatif dengan "/", misalnya biaya/uang-kuliah.md."""
    return path.relative_to(folder).as_posix()


def doc_id_for(folder: Path, path: Path) -> str:
    return path.relative_to(folder).with_suffix("").as_posix()


def _page_overrides(folder: Path, path: Path, raw: Any) -> dict[str, str]:
    """page_overrides dari sources.yaml: {"8-9": "_overrides/x.md"}; hanya untuk PDF."""
    if raw is None:
        return {}
    key = source_key(folder, path)
    if path.suffix.lower() != ".pdf":
        raise KbBuildError(ErrorCode.KB_INVALID, f"{key}: page_overrides hanya untuk PDF")
    if not isinstance(raw, dict):
        raise KbBuildError(ErrorCode.KB_INVALID, f"{key}: page_overrides harus mapping")
    out = {}
    for pages, md in raw.items():
        try:
            parse_page_range(pages)
        except ValueError as exc:
            raise KbBuildError(ErrorCode.KB_INVALID, f"{key}: {exc}") from exc
        if not str(md).lower().endswith(".md"):
            raise KbBuildError(ErrorCode.KB_INVALID, f"{key}: page_overrides harus file .md")
        out[str(pages)] = str(md)
    return out


# --- bagian yang digantikan dokumen lain (Langkah 64) ---
#
# sources.yaml: superseded_sections: [{section, superseded_by, reason}] pada entri dokumen apa pun.
# Saat build, chunk yang jalur bagiannya memuat bagian itu tidak dimasukkan ke indeks.

_SUPERSEDED_FIELDS = ("section", "superseded_by", "reason")


def _superseded_sections(folder: Path, path: Path, raw: Any) -> tuple[SupersededSection, ...]:
    if raw is None:
        return ()
    key = source_key(folder, path)
    if not isinstance(raw, list):
        raise KbBuildError(ErrorCode.KB_INVALID, f"{key}: superseded_sections harus daftar")
    out = []
    for item in raw:
        if not isinstance(item, dict):
            raise KbBuildError(
                ErrorCode.KB_INVALID, f"{key}: superseded_sections berisi {{section, ...}}"
            )
        values = {name: str(item.get(name) or "").strip() for name in _SUPERSEDED_FIELDS}
        empty = [name for name, value in values.items() if not value]
        if empty:
            raise KbBuildError(
                ErrorCode.KB_INVALID, f"{key}: superseded_sections tanpa field {empty}"
            )
        if any(s.section == values["section"] for s in out):
            raise KbBuildError(
                ErrorCode.KB_INVALID, f"{key}: superseded_sections {values['section']!r} ganda"
            )
        out.append(SupersededSection(**values))
    return tuple(out)


# label penomoran di akhir bagian: "4", "1.2", "II"; hanya label yang boleh dicocokkan sebagai awalan
_LABEL_END = re.compile(r"(?:\d+(?:\.\d+)*|[IVXLC]+)")


def _segment_matches(segment: str, part: str) -> bool:
    # "Pasal 4" cocok dengan "Pasal 4" dan "Pasal 4 Jalur ...", tidak dengan "Pasal 40 ...";
    # "Pasal" saja (tanpa nomor) hanya cocok dengan bagian bernama persis "Pasal"
    if segment == part:
        return True
    return _LABEL_END.fullmatch(part.split()[-1]) is not None and segment.startswith(part + " ")


def section_matches(section_path: str, section: str) -> bool:
    """True bila jalur bagian chunk memuat bagian itu, dicocokkan tepat per bagian.

    section boleh berupa jalur ("BAB II › Pasal 4"); bagian-bagiannya harus berurutan di jalur.
    """
    parts = [p.strip() for p in section.split(SECTION_SEP.strip()) if p.strip()]
    segments = [s.strip() for s in section_path.split(SECTION_SEP)] if section_path else []
    if not parts:
        return False
    for start in range(len(segments) - len(parts) + 1):
        window = segments[start : start + len(parts)]
        if all(_segment_matches(seg, part) for seg, part in zip(window, parts, strict=True)):
            return True
    return False


def parse_file(
    folder: Path, path: Path, page_overrides: Mapping[str, str] | None = None
) -> ParsedDocument | None:
    """Parse satu file dengan doc_id dan group sesuai posisinya di folder sumber.

    page_overrides: rentang halaman -> path .md relatif di folder (hanya PDF).
    """
    overrides = dict(page_overrides or {})
    if overrides:
        resolved = {parse_page_range(k): folder / v for k, v in overrides.items()}
        parsed = PdfParser().parse(path, resolved)
    else:
        parsed = PARSERS[path.suffix.lower()].parse(path)
    if parsed is None:
        return None
    parts = path.relative_to(folder).parts
    return dataclasses.replace(
        parsed,
        doc_id=doc_id_for(folder, path),
        group=parts[0] if len(parts) > 1 else "",
        page_overrides=MappingProxyType(overrides),
    )


def load_sources(
    folder: Path, source_type: SourceType, allow_missing_date: bool = False
) -> list[ParsedDocument]:
    folder = Path(folder)
    if source_type not in ("official", "dev"):
        raise ValueError(f"source_type tidak dikenal: {source_type!r}")
    sources = _read_sources(folder)

    documents: list[ParsedDocument] = []
    missing_date: list[str] = []
    seen_ids: dict[str, str] = {}
    for path in iter_source_files(folder):
        key = source_key(folder, path)
        doc_id = doc_id_for(folder, path)
        if doc_id in seen_ids:
            raise KbBuildError(
                ErrorCode.KB_INVALID, f"doc_id {doc_id!r} ganda: {seen_ids[doc_id]} dan {key}"
            )
        seen_ids[doc_id] = key

        meta = sources.get(key, {})
        overrides = _page_overrides(folder, path, meta.get("page_overrides"))
        parsed = parse_file(folder, path, overrides)
        if parsed is None:
            continue
        valid_from = _parse_date(meta.get("valid_from"))
        if valid_from is None:
            missing_date.append(key)
        documents.append(
            dataclasses.replace(
                parsed,
                doc_title=str(meta.get("title") or path.stem),
                valid_from=valid_from,
                source_type=source_type,
                superseded_sections=_superseded_sections(
                    folder, path, meta.get("superseded_sections")
                ),
            )
        )

    if missing_date and not allow_missing_date:
        raise KbBuildError(
            ErrorCode.KB_MISSING_DATE,
            f"valid_from belum diisi di {SOURCES_FILE} untuk: {', '.join(missing_date)}",
        )
    return documents
